"""Protected, reviewed whole-workbook replacement; transfers remain session-local."""
from __future__ import annotations
import base64
import hashlib
from io import BytesIO
from zipfile import ZipFile
import openpyxl
import streamlit as st
import admin_auth
import public_workbook
import race_github

LIMIT = 10 * 1024 * 1024
KEY = 'race_import_excel_upload'

def validate(content: bytes) -> None:
    public_workbook.validate_snapshot(content)
    with ZipFile(BytesIO(content)) as archive:
        if any(name.lower().startswith(('xl/externallinks/', 'xl/embeddings/')) or 'vbaproject' in name.lower() for name in archive.namelist()):
            raise ValueError('Macros, embedded objects and external workbook links are not accepted.')

def compare(previous: bytes, candidate: bytes) -> dict:
    validate(candidate)
    books = [openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=False) for data in (previous, candidate)]
    try:
        before, after = books
        missing = set(before.sheetnames) - set(after.sheetnames)
        if missing:
            raise ValueError('The uploaded Excel is missing existing worksheets: ' + ', '.join(sorted(missing)))
        summary=[]; sample=[]
        from itertools import zip_longest
        for sheet in after:
            old = before[sheet.title].iter_rows(values_only=True) if sheet.title in before.sheetnames else []
            changed=deleted=0
            for rowno, (left,right) in enumerate(zip_longest(old,sheet.iter_rows(values_only=True),fillvalue=()),1):
                for colno,(a,b) in enumerate(zip_longest(left,right,fillvalue=None),1):
                    if a == b: continue
                    changed+=1
                    if a is not None and b is None: deleted+=1
                    if len(sample)<100:
                        sample.append({'Sheet':sheet.title,'Cell':f'{openpyxl.utils.get_column_letter(colno)}{rowno}','Before':str(a) if a is not None else '', 'After':str(b) if b is not None else ''})
            summary.append({'Sheet':sheet.title,'Changed cells':changed,'Cleared cells':deleted})
        return {'summary':summary,'sample':sample,'digest':hashlib.sha256(candidate).hexdigest()}
    finally:
        for book in books: book.close()

def publish(config, content: bytes, *, expected_sha: str, approved_digest: str, approved: bool):
    if not admin_auth.is_current_admin() or not approved:
        raise PermissionError('Admin sign-in and explicit approval are required.')
    if hashlib.sha256(content).hexdigest() != approved_digest:
        raise ValueError('The approved Excel has changed. Review it again.')
    client = race_github.GitHubAppClient(config)
    token = client.create_installation_token()
    remote = client._fetch_with_token(token)
    if remote.blob_sha != expected_sha:
        raise race_github.GitHubConflictError('GitHub changed after preview. Download the latest Excel and review again.')
    compare(remote.content,content)
    if remote.content == content:
        raise ValueError('This Excel is already published.')
    return client._publish_updated_workbook(token=token,remote=remote,updated_bytes=content,
        message='Admin: publish reviewed Excel adjustments',operation='publish reviewed Excel')

UPLOAD_JS = r"""
export default function({parentElement,data,setStateValue}) {
 const input=parentElement.querySelector('input'),label=parentElement.querySelector('label'),status=parentElement.querySelector('p');
 label.firstChild.textContent=data.label; input.setAttribute('aria-label',data.label);
 input.onchange=async()=>{setStateValue('file',null);const file=input.files[0];if(!file)return;
   if(!/\.xlsx$/i.test(file.name)||file.size>10485760||!file.size){status.textContent=data.invalid;return;}
   try {const bytes=new Uint8Array(await file.arrayBuffer());let binary='';for(let i=0;i<bytes.length;i+=32768)binary+=String.fromCharCode(...bytes.subarray(i,i+32768));
    if(input.files[0]!==file)return;
    setStateValue('file',{name:file.name,content:btoa(binary)});status.textContent=file.name;
   }catch(_){status.textContent=data.invalid;}
 };
}
"""

def render(config, *, lang: str) -> None:
    if not admin_auth.is_current_admin():
        st.stop()
    pt=lang=='pt'
    st.subheader('Upload latest Excel' if not pt else 'Carregar Excel mais recente')
    st.caption('Escolhe um ficheiro .xlsx (até 10 MB), revê as alterações e aprova a publicação. A versão anterior fica no histórico do GitHub.' if pt else
               'Choose an .xlsx file (up to 10 MB), review the changes, then approve publication. The previous version remains in GitHub history.')
    component=st.components.v2.component('f1_excel_upload',html='<label>Excel <input type="file" accept=".xlsx" /></label><p role="status"></p>',js=UPLOAD_JS,
        css='label,p {font:14px system-ui;color:#fafafa} input {display:block;margin:12px 0} input::file-selector-button {padding:12px;color:white;background:#20232c;border:1px solid #555;border-radius:8px}')
    result=component(key=KEY,data={'label':'Escolher Excel' if pt else 'Choose Excel','invalid':'Ficheiro .xlsx inválido (máx. 10 MB).' if pt else 'Choose a valid .xlsx file (max. 10 MB).'},on_file_change=lambda:st.session_state.pop(KEY+'_preview',None),height='content')
    payload = result.get('file') if hasattr(result,'get') else None
    if not payload: return
    try:
        encoded=payload.get('content','')
        if not isinstance(encoded,str) or len(encoded)>LIMIT*4//3+4: raise ValueError('Excel exceeds 10 MB.')
        content=base64.b64decode(encoded,validate=True)
        if not isinstance(payload.get('name'),str) or not payload['name'].lower().endswith('.xlsx'): raise ValueError('Choose an .xlsx file.')
        if st.button('Pré-visualizar alterações' if pt else 'Preview Excel changes'):
            remote=race_github.fetch_remote_workbook(config)
            preview=compare(remote.content,content)
            preview['sha']=remote.blob_sha
            st.session_state[KEY+'_preview']=preview
        preview=st.session_state.get(KEY+'_preview')
        if not preview or preview['digest']!=hashlib.sha256(content).hexdigest(): return
        st.table(preview['summary'])
        st.caption('Primeiras 100 alterações de células; estilos e fórmulas do ficheiro carregado também serão substituídos.' if pt else 'First 100 cell changes; uploaded formatting and formulas will also replace the current file.')
        st.table(preview['sample'])
        approved=st.checkbox('Aprovo substituir o Excel completo no GitHub, incluindo as células removidas acima.' if pt else 'I approve replacing the entire GitHub Excel, including the cleared cells shown above.',key=KEY+'_approve_'+preview['sha']+'_'+preview['digest'])
        if st.button('Publicar Excel aprovado' if pt else 'Publish approved Excel',disabled=not approved):
            # Consume the preview before any uncertain network response; never silently retry a write.
            st.session_state.pop(KEY+'_preview',None)
            _,url,_=publish(config,content,expected_sha=preview['sha'],approved_digest=preview['digest'],approved=approved)
            admin_auth.clear_race_import_state(st.session_state)
            st.success('Excel publicado. O Dashboard irá carregar esta versão.' if pt else 'Excel published. The Dashboard will load this version.')
            st.link_button('GitHub commit',url)
            st.stop()
    except (ValueError,TypeError,KeyError,race_github.GitHubPersistenceError,OSError) as exc:
        st.session_state.pop(KEY+'_preview',None)
        st.error(('Não foi possível publicar. Descarrega o Excel mais recente e revê novamente.' if pt else 'Unable to publish. Download the latest Excel and review again.') + (' ' + str(exc) if isinstance(exc,ValueError) else ''))
